"""Proje klasörü depolama katmanı: Google Drive API v3 veya yerel klasör.

Google Drive iki şekilde bağlanabilir:

* **Hizmet hesabı (service_account):** Google Cloud'da bir hizmet hesabı açılır, JSON
  anahtarı yüklenir ve hizmet hesabının e-postası bir **Paylaşılan Drive**'a (Shared
  Drive) "İçerik yöneticisi" olarak eklenir. Hizmet hesaplarının kendi depolama kotası
  olmadığı için klasör ID'si mutlaka bir Paylaşılan Drive içinde olmalıdır.
* **Kullanıcı hesabı (oauth):** OAuth istemci dosyası yüklenir ve bir kez
  `python manage.py drive-yetkilendir` çalıştırılarak token alınır. Dosyalar o
  kullanıcının Drive'ında oluşur.

Drive kapalıysa aynı klasör yapısı data/projeler/ altında yerel olarak oluşturulur.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from . import config, settings

FOLDER_MIME = "application/vnd.google-apps.folder"
GDOC_MIME = "application/vnd.google-apps.document"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SCOPES = ["https://www.googleapis.com/auth/drive"]
TOKEN_PATH = config.SECRETS_DIR / "drive_token.json"


class LocalStorage:
    name = "yerel"

    def create_folder(self, name: str, parent: str | None = None) -> dict:
        base = Path(parent) if parent else config.PROJECT_DIR
        path = base / name
        path.mkdir(parents=True, exist_ok=True)
        return {"id": str(path), "url": path.resolve().as_uri()}

    def upload(self, local_path: str, folder_id: str, convert: bool = False, name: str | None = None) -> dict:
        src = Path(local_path)
        dest = Path(folder_id) / (name or src.name)
        if src.resolve() != dest.resolve():
            shutil.copy2(src, dest)
        return {"id": str(dest), "url": dest.resolve().as_uri()}

    def share(self, file_id: str, email: str, role: str = "writer") -> None:
        return None

    def export_text(self, file_id: str) -> str:
        return docx_text(file_id)


class GoogleDriveStorage:
    name = "google_drive"

    def __init__(self) -> None:
        from googleapiclient.discovery import build  # type: ignore

        self.service = build("drive", "v3", credentials=load_credentials(), cache_discovery=False)
        self.root = str(settings.get("drive_ana_klasor_id") or "") or None

    def create_folder(self, name: str, parent: str | None = None) -> dict:
        body = {"name": name, "mimeType": FOLDER_MIME}
        parent = parent or self.root
        if parent:
            body["parents"] = [parent]
        f = self.service.files().create(body=body, fields="id, webViewLink", supportsAllDrives=True).execute()
        return {"id": f["id"], "url": f.get("webViewLink")}

    def upload(self, local_path: str, folder_id: str, convert: bool = False, name: str | None = None) -> dict:
        from googleapiclient.http import MediaFileUpload  # type: ignore

        src = Path(local_path)
        body = {"name": name or src.name, "parents": [folder_id]}
        mimetype = None
        if src.suffix.lower() == ".docx":
            mimetype = DOCX_MIME
            if convert:
                body["mimeType"] = GDOC_MIME
                body["name"] = (name or src.stem).rsplit(".docx", 1)[0]
        elif src.suffix.lower() == ".pdf":
            mimetype = "application/pdf"
        media = MediaFileUpload(str(src), mimetype=mimetype, resumable=False)
        f = self.service.files().create(body=body, media_body=media, fields="id, webViewLink",
                                        supportsAllDrives=True).execute()
        return {"id": f["id"], "url": f.get("webViewLink")}

    def share(self, file_id: str, email: str, role: str = "writer") -> None:
        self.service.permissions().create(
            fileId=file_id, body={"type": "user", "role": role, "emailAddress": email},
            sendNotificationEmail=settings.get_bool("drive_paylasim_bildirimi"), supportsAllDrives=True,
        ).execute()

    def export_text(self, file_id: str) -> str:
        meta = self.service.files().get(fileId=file_id, fields="mimeType", supportsAllDrives=True).execute()
        if meta.get("mimeType") == GDOC_MIME:
            data = self.service.files().export(fileId=file_id, mimeType="text/plain").execute()
            return data.decode("utf-8", errors="ignore") if isinstance(data, bytes) else str(data)
        import io
        data = self.service.files().get_media(fileId=file_id, supportsAllDrives=True).execute()
        return docx_text(io.BytesIO(data))

    def whoami(self) -> str:
        about = self.service.about().get(fields="user(emailAddress,displayName)").execute()
        return about.get("user", {}).get("emailAddress", "?")


def load_credentials():
    mode = str(settings.get("drive_mod") or "service_account")
    if mode == "oauth":
        from google.oauth2.credentials import Credentials  # type: ignore
        from google.auth.transport.requests import Request  # type: ignore

        if not TOKEN_PATH.exists():
            raise RuntimeError("Drive OAuth token yok. Sunucuda `python manage.py drive-yetkilendir` çalıştırın.")
        creds = Credentials.from_authorized_user_file(str(TOKEN_PATH), SCOPES)
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())
            TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")
        return creds
    from google.oauth2 import service_account  # type: ignore

    path = str(settings.get("drive_kimlik_dosyasi") or "")
    if not path or not Path(path).exists():
        raise RuntimeError("Hizmet hesabı JSON dosyası bulunamadı. Yönetim Paneli > Entegrasyonlar'dan yükleyin.")
    return service_account.Credentials.from_service_account_file(path, scopes=SCOPES)


def oauth_authorize(client_secret_path: str) -> str:
    """Komut satırından bir kez çalıştırılır; token dosyasını kaydeder."""
    from google_auth_oauthlib.flow import InstalledAppFlow  # type: ignore

    flow = InstalledAppFlow.from_client_secrets_file(client_secret_path, SCOPES)
    creds = flow.run_local_server(port=0, open_browser=True)
    TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")
    return str(TOKEN_PATH)


def get_storage():
    if settings.get_bool("drive_aktif"):
        return GoogleDriveStorage()
    return LocalStorage()


def docx_text(path_or_stream) -> str:
    from docx import Document  # type: ignore

    doc = Document(path_or_stream)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_connection() -> str:
    storage = GoogleDriveStorage()
    who = storage.whoami()
    if storage.root:
        storage.service.files().get(fileId=storage.root, fields="id,name", supportsAllDrives=True).execute()
    return who
