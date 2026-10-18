"""Varsayılan Word (.docx) şablonlarını üretir.

Şablonlar docxtpl (Jinja2) yer tutucuları içerir: {{ baslik }}, {{ ozet }} ...
Kurumlar kendi şablonlarını Yönetim Paneli > Şablonlar bölümünden yükleyebilir;
kullanılabilecek değişkenlerin listesi aynı ekranda gösterilir.

Çalıştırma:  python tools/build_templates.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "templates"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x55, 0x5F, 0x6D)
HEAD_FILL = "1F3864"
LIGHT_FILL = "EEF2F8"
NOTE_FILL = "FFF8E1"


# ---------------------------------------------------------------------------
# Yardımcılar
# ---------------------------------------------------------------------------
def new_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15
    return doc


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para(doc_or_cell, text: str = "", bold=False, size=None, color=None, align=None, italic=False, space_after=None):
    p = doc_or_cell.add_paragraph()
    if text:
        add_run(p, text, bold=bold, size=size, color=color, italic=italic)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_run(p, text, bold=False, size=None, color=None, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_run(p, text, bold=True, size=12.5 if level == 1 else 11, color=NAVY)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", HEAD_FILL)):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)
    return p


def hint(doc, text: str):
    return para(doc, text, italic=True, size=9, color=GREY)


def var(doc, name: str):
    """Tek başına bir değişken paragrafı."""
    return para(doc, "{{ " + name + " }}")


def note_box(doc, text: str):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    set_widths(t, [17])
    c = t.rows[0].cells[0]
    shade(c, NOTE_FILL)
    c.paragraphs[0].text = ""
    add_run(c.paragraphs[0], text, size=9)
    doc.add_paragraph()
    return t


def kv_table(doc, rows: list[tuple[str, str]], widths=(5.5, 11.5)):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        a, b = t.rows[i].cells
        shade(a, LIGHT_FILL)
        a.paragraphs[0].text = ""
        add_run(a.paragraphs[0], k, bold=True, size=9.5)
        b.paragraphs[0].text = ""
        add_run(b.paragraphs[0], v, size=10)
    set_widths(t, widths)
    doc.add_paragraph()
    return t


def loop_table(doc, headers: list[str], loop_var: str, items: str, cells: list[str], widths: list[float] | None = None,
               font_size: float = 9, footer: list[str] | None = None):
    """docxtpl satır döngülü tablo: başlık, {%tr for%}, veri satırı, {%tr endfor%}."""
    n = len(headers)
    t = doc.add_table(rows=4 + (1 if footer else 0), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, HEAD_FILL)
        c.paragraphs[0].text = ""
        add_run(c.paragraphs[0], h, bold=True, size=font_size, color=RGBColor(0xFF, 0xFF, 0xFF))
    t.rows[1].cells[0].paragraphs[0].text = "{%tr for " + loop_var + " in " + items + " %}"
    for i, expr in enumerate(cells):
        c = t.rows[2].cells[i]
        c.paragraphs[0].text = ""
        add_run(c.paragraphs[0], "{{ " + expr + " }}", size=font_size)
    t.rows[3].cells[0].paragraphs[0].text = "{%tr endfor %}"
    if footer:
        for i, f in enumerate(footer):
            c = t.rows[4].cells[i]
            shade(c, LIGHT_FILL)
            c.paragraphs[0].text = ""
            add_run(c.paragraphs[0], f, bold=True, size=font_size)
    if widths:
        set_widths(t, widths)
    doc.add_paragraph()
    return t


def set_widths(t, widths) -> None:
    """Sütun genişliklerini hem tblGrid hem hücre düzeyinde ayarlar (Word ve LibreOffice uyumlu)."""
    t.autofit = False
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
    for row in t.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def title_block(doc, line1: str, line2: str, line3: str | None = None):
    para(doc, "{{ universite_adi }}", bold=True, size=10, color=GREY, align="center", space_after=0)
    para(doc, "{{ kurum_adi }}", bold=True, size=10, color=GREY, align="center", space_after=10)
    para(doc, line1, bold=True, size=14, color=NAVY, align="center", space_after=2)
    para(doc, line2, bold=True, size=11.5, align="center", space_after=2)
    if line3:
        para(doc, line3, italic=True, size=9.5, color=GREY, align="center", space_after=8)


def footer_text(doc, text: str):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.text = ""
    add_run(p, text, size=8, color=GREY)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def gantt_table(doc):
    headers = ["No"] + [str(i) for i in range(1, 13)]
    cells = ["g.no"] + [f"g.a{i}" for i in range(1, 13)]
    hint(doc, "Zaman çizelgesi (ay bazında; ■ = iş adımının sürdüğü aylar)")
    loop_table(doc, headers, "g", "gantt", cells, widths=[1.2] + [1.3] * 12, font_size=8.5)


def team_table(doc):
    loop_table(doc, ["Ad Soyad", "Rol", "Bölüm / Kurum", "Yetkinlikler", "ARBİS"], "u", "ekip",
               ["u.ad", "u.rol", "u.bolum", "u.yetkinlik", "u.arbis"], widths=[3.6, 2.6, 3.6, 5.6, 1.6])


def control_table(doc):
    loop_table(doc, ["Durum", "Alan", "Açıklama"], "k", "kontroller", ["k.durum", "k.alan", "k.mesaj"],
               widths=[2.2, 5, 9.8], font_size=8.5)


def budget_table(doc):
    loop_table(doc, ["Kalem", "Tür", "Adet", "Birim Fiyat (TL)", "Tutar (TL)", "Gerekçe"], "b", "butce",
               ["b.kalem", "b.tur", "b.adet", "b.birim", "b.tutar", "b.gerekce"],
               widths=[3.8, 2.6, 1.3, 2.4, 2.4, 4.5], footer=["TOPLAM", "", "", "", "{{ butce_toplam }}", "Üst limit: {{ butce_limit }} TL"])


def work_plan_table(doc):
    loop_table(doc, ["No", "İş Adımı / Ana Faaliyet", "Sorumlu", "Ay", "Başarı Ölçütü", "Katkı (%)"], "r", "is_plani",
               ["r.no", "r.is_paketi", "r.sorumlu", "r.aylar", "r.basari_olcutu", "r.katki"],
               widths=[1, 4.6, 2.8, 1.6, 5.6, 1.4], footer=["", "TOPLAM", "", "", "", "{{ katki_toplam }}"])


def risk_table(doc):
    loop_table(doc, ["İş Adımı No", "Önemli Riskler", "Risk Yönetimi (B Planı)"], "r", "risk_plani",
               ["r.no", "r.risk", "r.b_plani"], widths=[2.2, 6.8, 8])


def section(doc, title: str, var_name: str, hint_text: str | None = None, level: int = 2):
    heading(doc, title, level)
    if hint_text:
        hint(doc, hint_text)
    var(doc, var_name)


# ---------------------------------------------------------------------------
# 2209-A / 2209-B TYBS hazırlık taslağı
# ---------------------------------------------------------------------------
def build_2209(tur: str) -> Document:
    doc = new_doc()
    if tur == "A":
        prog = "TÜBİTAK 2209-A ÜNİVERSİTE ÖĞRENCİLERİ ARAŞTIRMA PROJELERİ DESTEKLEME PROGRAMI"
    else:
        prog = "TÜBİTAK 2209-B ÜNİVERSİTE ÖĞRENCİLERİ SANAYİYE YÖNELİK ARAŞTIRMA PROJELERİ DESTEĞİ PROGRAMI"
    title_block(doc, prog, "Araştırma Önerisi — TYBS Hazırlık Taslağı ({{ donem }} Çağrısı)",
                "Proje kodu: {{ proje_kodu }} · Oluşturma: {{ tarih }}")
    note_box(doc, "Bu belge TYBS (tybs.tubitak.gov.tr) başvuru ekranlarıyla aynı başlık sırasında hazırlanmıştır. "
                  "Araştırma önerisi TYBS'ye girilen bilgilerden otomatik oluşturulur; bu dosya TÜBİTAK'a yüklenmez. "
                  "Ekip, içeriği bu taslakta birlikte yazar ve ilgili ekranlara aktarır. "
                  "[DOLDURULACAK] ile işaretli alanlar tamamlanmadan başvuru gönderilmemelidir.")
    heading(doc, "GENEL BİLGİLER")
    rows = [
        ("Araştırma Önerisinin Başlığı", "{{ baslik }}"),
        ("Proje Yürütücüsü", "{{ yurutucu }}"),
        ("Proje Ortakları", "{{ ortaklar }}"),
        ("Akademik Danışman", "{{ danisman }}"),
    ]
    if tur == "B":
        rows += [("Sanayi Kuruluşu", "{{ sanayi_kurulus }}"), ("Sanayi Danışmanı", "{{ sanayi_danismani }}")]
    rows += [
        ("Araştırmanın Yürütüleceği Kurum", "{{ kurum_adi }} — {{ bolum }}"),
        ("Proje Süresi", "{{ sure_ay }} ay (en çok 12 ay)"),
        ("Talep Edilen Bütçe", "{{ butce_toplam }} TL (üst limit {{ butce_limit }} TL)"),
        ("Anahtar Kelimeler", "{{ anahtar_kelimeler }}"),
        ("Etik Kurul Gereksinimi", "{{ etik_kurul }}"),
        ("Önerilen TEKNOFEST Çapraz Başvurusu", "{{ teknofest_onerisi }}"),
    ]
    kv_table(doc, rows)

    heading(doc, "ÖZET")
    hint(doc, "25–450 kelime. Bilimsel nitelik, yöntem, proje yönetimi ve yaygın etkiyi kapsamalı; en son yazılması önerilir. "
              "Mevcut uzunluk: {{ ozet_kelime }} kelime.")
    var(doc, "ozet")

    heading(doc, "1. ARAŞTIRMA ÖNERİSİNİN BİLİMSEL NİTELİĞİ")
    if tur == "A":
        section(doc, "1.1. Konunun Önemi ve Araştırma Önerisinin Bilimsel Niteliği", "konu_onemi",
                "En az 350 karakter. Konunun kapsamı ve önemi, literatürdeki boşluk/sorun, araştırma sorusu ve varsa hipotez. "
                "12. Kalkınma Planı veya 2030 Sanayi ve Teknoloji Stratejisi ile ilişkisi varsa belirtin.")
        section(doc, "1.2. Amaç ve Hedefler", "amac_hedefler",
                "En az 350 karakter. Açık, ölçülebilir, gerçekçi ve ulaşılabilir hedefler.")
    else:
        section(doc, "1.1. Amaç ve Hedefler", "amac_hedefler",
                "En az 350 karakter. Açık, ölçülebilir, gerçekçi ve ulaşılabilir hedefler; sanayi probleminin tanımı.")
        section(doc, "1.2. Yenilikçi Yönü ve Teknolojik Değeri", "yenilikci_yon",
                "En az 350 karakter. Mevcut çözümlere göre yenilik, teknolojik katkı ve sanayiye sağlayacağı değer.")

    heading(doc, "2. YÖNTEM")
    hint(doc, "En az 350 karakter. Araştırma tasarımı, değişkenler, ölçüm/analiz yöntemleri, varsa ön çalışma; yöntem çalışma takvimiyle ilişkilendirilmeli.")
    var(doc, "yontem")

    heading(doc, "3. PROJE YÖNETİMİ")
    heading(doc, "3.1. Çalışma Takvimi", 2)
    hint(doc, "Başarı ölçütleri ölçülebilir olmalı ve katkı sütununun toplamı %100 olmalıdır. Literatür taraması, rapor yazımı, "
              "sonuçların yaygınlaştırılması ve malzeme alımı ayrı iş adımı olarak yazılmaz.")
    work_plan_table(doc)
    gantt_table(doc)
    heading(doc, "3.2. Risk Yönetimi", 2)
    hint(doc, "B planları ana amaçtan sapmaya yol açmamalıdır; yöntem değişiyorsa ayrıntılandırılmalıdır.")
    risk_table(doc)
    section(doc, "3.3. Araştırma Olanakları", "arastirma_olanaklari",
            "Kurumda bulunan ve projede kullanılacak altyapı, laboratuvar, atölye, cihaz ve makineler.")

    if tur == "A":
        section(doc, "4. ARAŞTIRMA ÖNERİSİNİN YAYGIN ETKİSİ", "yaygin_etki",
                "Beklenen çıktılar amaçlarına göre gruplanmalı ve ölçülebilir, gerçekçi hedeflere dayanmalı.", level=1)
    else:
        section(doc, "4. SANAYİ ODAKLI ÇIKTILARI VE YAYGIN ETKİ", "yaygin_etki",
                "Sanayi kuruluşuna sağlanacak çıktılar (ürün, prototip, süreç iyileştirme vb.) ölçülebilir hedeflerle.", level=1)
    section(doc, "5. BELİRTMEK İSTEDİĞİNİZ DİĞER KONULAR", "diger_konular",
            "İsteğe bağlı; yalnızca değerlendirmeye katkı sağlayacak bilgi (grafik, tablo vb.).", level=1)
    heading(doc, "6. EKLER")
    hint(doc, "Varsa ölçek, anket formu vb. destekleyici materyaller TYBS'deki ilgili adıma eklenir.")
    heading(doc, "KAYNAKLAR")
    hint(doc, "TÜBİTAK kaynakça kılavuzuna uygun yazılmalı; her kaynak metinde atıf almalı; DOI'si olan kaynakta DOI zorunludur.")
    var(doc, "kaynaklar")

    heading(doc, "EK-A: BÜTÇE TALEBİ (TYBS bütçe ekranı için hazırlık)")
    budget_table(doc)
    heading(doc, "EK-B: EKİP VE GÖREV DAĞILIMI")
    team_table(doc)
    heading(doc, "EK-C: OTOMATİK ÖN KONTROL SONUCU ({{ tarih }})")
    hint(doc, "Sistem tarafından TYBS kurallarına göre yapılan kontrol. 'HATA' satırları başvurudan önce giderilmelidir.")
    control_table(doc)
    footer_text(doc, "{{ kurum_adi }} · Akıllı Proje Kuluçka ve Otomasyon Sistemi · {{ tarih }}")
    return doc


# ---------------------------------------------------------------------------
# TEKNOFEST Ön Değerlendirme Raporu taslağı
# ---------------------------------------------------------------------------
def build_odr() -> Document:
    doc = new_doc()
    title_block(doc, "TEKNOFEST {{ sezon }} — {{ yarisma_adi }}", "Ön Değerlendirme / Proje Detay Raporu — Hazırlık Taslağı",
                "Proje kodu: {{ proje_kodu }} · Oluşturma: {{ tarih }}")
    note_box(doc, "Rapor başlıkları ve sayfa sınırları yarışmaya göre değişir. Yarışmanın resmi rapor şablonu ve şartnamesi "
                  "proje klasöründeki TEKNOFEST ve Şartname alt klasörlerine indirildi (yayımlanmışsa). Bu taslaktaki içerik "
                  "resmi şablona aktarılarak teslim edilmelidir.")
    heading(doc, "TAKIM VE YARIŞMA BİLGİLERİ")
    kv_table(doc, [
        ("Proje Adı", "{{ baslik }}"),
        ("Takım Adı", "{{ takim_adi }}"),
        ("Takım ID / Başvuru ID", "[DOLDURULACAK: TEKNOFEST başvurusundan sonra yazılacak]"),
        ("Yarışma", "{{ yarisma_adi }}"),
        ("Yarışma Sayfası", "{{ yarisma_url }}"),
        ("Son Başvuru Tarihi", "{{ son_basvuru }}"),
        ("Katılım Seviyesi", "{{ katilim_seviyesi }}"),
        ("Proje–Yarışma Uyum Skoru", "%{{ uyum_skoru }}"),
        ("Danışman", "{{ danisman }}"),
        ("Kurum", "{{ universite_adi }} {{ kurum_adi }}"),
    ])
    heading(doc, "TAKIM ÜYELERİ")
    team_table(doc)
    section(doc, "1. Proje Özeti", "ozet", "Projenin tamamını kısaca anlatan özet.", level=1)
    section(doc, "2. Problem / Sorun", "problem", "Çözülmek istenen sorun ve önemi.", level=1)
    section(doc, "3. Çözüm Önerisi", "cozum", "Önerilen çözüm ve hedefler.", level=1)
    section(doc, "4. Yöntem", "yontem", "Kullanılacak yöntem, teknoloji ve sistem mimarisi.", level=1)
    section(doc, "5. Yenilikçi (Özgün) Yönü", "yenilikci_yon", "Benzer çözümlerden farkı, özgün katkı.", level=1)
    section(doc, "6. Uygulanabilirlik ve Ticarileşme Potansiyeli", "uygulanabilirlik",
            "Gerçek hayatta uygulanabilirlik, ürünleşme/ticarileşme potansiyeli.", level=1)
    heading(doc, "7. Tahmini Maliyet ve Proje Zaman Planlaması")
    budget_table(doc)
    work_plan_table(doc)
    section(doc, "8. Projenin Hedef Kitlesi (Kullanıcılar)", "hedef_kitle", "Projeden fayda sağlayacak kullanıcılar.", level=1)
    heading(doc, "9. Riskler")
    risk_table(doc)
    section(doc, "10. Kaynaklar", "kaynaklar", None, level=1)
    footer_text(doc, "{{ kurum_adi }} · Akıllı Proje Kuluçka ve Otomasyon Sistemi · {{ tarih }}")
    return doc


# ---------------------------------------------------------------------------
# Bitirme projesi öneri formu
# ---------------------------------------------------------------------------
def build_bitirme() -> Document:
    doc = new_doc()
    title_block(doc, "BİTİRME PROJESİ ÖNERİ FORMU", "{{ donem }} Dönemi", "Proje kodu: {{ proje_kodu }} · Oluşturma: {{ tarih }}")
    heading(doc, "GENEL BİLGİLER")
    kv_table(doc, [
        ("Proje Adı", "{{ baslik }}"),
        ("Program / Bölüm", "{{ bolum }}"),
        ("Danışman", "{{ danisman }}"),
        ("Sanayi Paydaşı (varsa)", "{{ sanayi_kurulus }}"),
        ("Proje Süresi", "{{ sure_ay }} ay"),
        ("Anahtar Kelimeler", "{{ anahtar_kelimeler }}"),
    ])
    heading(doc, "PROJE EKİBİ")
    team_table(doc)
    section(doc, "1. Özet", "ozet", None, level=1)
    section(doc, "2. Amaç ve Hedefler", "amac_hedefler", None, level=1)
    section(doc, "3. Yöntem", "yontem", None, level=1)
    section(doc, "4. Beklenen Çıktılar", "yaygin_etki", None, level=1)
    heading(doc, "5. İş-Zaman Planı")
    work_plan_table(doc)
    gantt_table(doc)
    heading(doc, "6. Malzeme ve Maliyet")
    budget_table(doc)
    section(doc, "7. Kaynaklar", "kaynaklar", None, level=1)
    heading(doc, "ONAY")
    t = doc.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["Öğrenci(ler)", "Danışman", "Bölüm Başkanı"]):
        c = t.rows[0].cells[i]
        shade(c, LIGHT_FILL)
        c.paragraphs[0].text = ""
        add_run(c.paragraphs[0], h, bold=True, size=9.5)
        t.rows[1].cells[i].paragraphs[0].text = "\n\nİmza / Tarih"
    footer_text(doc, "{{ kurum_adi }} · Akıllı Proje Kuluçka ve Otomasyon Sistemi · {{ tarih }}")
    return doc


# ---------------------------------------------------------------------------
# Başvuru kontrol listesi (her onaylanan projeye eklenir)
# ---------------------------------------------------------------------------
def build_checklist() -> Document:
    doc = new_doc()
    title_block(doc, "BAŞVURU KONTROL LİSTESİ VE REHBER", "{{ baslik }}", "{{ tur }} · Proje kodu: {{ proje_kodu }} · {{ tarih }}")
    heading(doc, "PROJE KLASÖRÜ")
    kv_table(doc, [
        ("Klasör bağlantısı", "{{ klasor_url }}"),
        ("Danışman", "{{ danisman }}"),
        ("Proje yürütücüsü / kaptan", "{{ yurutucu }}"),
        ("Ekip", "{{ ortaklar }}"),
    ])
    heading(doc, "YAKLAŞAN SON TARİHLER")
    loop_table(doc, ["Son Tarih", "Aşama", "Kalan Gün"], "t", "takvim", ["t.tarih", "t.baslik", "t.kalan"], widths=[3.5, 10.5, 3])
    heading(doc, "YAPILACAKLAR")
    loop_table(doc, ["✓", "Adım", "Sorumlu"], "a", "adimlar", ["a.kutu", "a.adim", "a.sorumlu"], widths=[1.2, 12, 3.8])
    heading(doc, "BAŞVURU REHBERİ")
    para(doc, "{%p for satir in rehber_satirlari %}")
    para(doc, "{{ satir }}")
    para(doc, "{%p endfor %}")
    heading(doc, "OTOMATİK ÖN KONTROL")
    control_table(doc)
    footer_text(doc, "{{ kurum_adi }} · Akıllı Proje Kuluçka ve Otomasyon Sistemi · {{ tarih }}")
    return doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_2209("A").save(OUT / "tubitak_2209a_sablon.docx")
    build_2209("B").save(OUT / "tubitak_2209b_sablon.docx")
    build_odr().save(OUT / "teknofest_odr_sablon.docx")
    build_bitirme().save(OUT / "bitirme_oneri_sablon.docx")
    build_checklist().save(OUT / "kontrol_listesi_sablon.docx")
    print("Şablonlar oluşturuldu:", OUT)


if __name__ == "__main__":
    sys.exit(main())
